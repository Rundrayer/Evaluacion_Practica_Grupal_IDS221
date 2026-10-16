

import argparse
import base64
import hashlib
import json
from datetime import datetime
from pathlib import Path

import matplotlib
matplotlib.use("Agg")                       
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from jinja2 import Environment, FileSystemLoader

AQUI = Path(__file__).resolve().parent                      

RAIZ = AQUI.parent                                          

CSV_DEFECTO = RAIZ / "Datos" / "Datasets" / "operacion_sedes.csv"
COLUMNAS = ["fecha", "sede_id", "region", "generacion_kwh", "demanda_kwh", "evento_programado"]


DEFINICION_DEFICIT = "suma por sede-día de max(demanda_kwh - generacion_kwh, 0)"


def sha256_archivo(ruta):
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        for bloque in iter(lambda: f.read(1024 * 1024), b""):
            h.update(bloque)
    return h.hexdigest()


def cargar(ruta_csv):
    df = pd.read_csv(ruta_csv, parse_dates=["fecha"])
    faltan = [c for c in COLUMNAS if c not in df.columns]
    if faltan:
        raise ValueError(f"Faltan columnas en el CSV: {faltan}")
    if df[COLUMNAS].isna().any().any():
        raise ValueError("El CSV tiene valores nulos en columnas requeridas")
    if (df[["generacion_kwh", "demanda_kwh"]] < 0).any().any():
        raise ValueError("Hay valores negativos de generacion o demanda")
    if df.duplicated(["fecha", "sede_id"]).any():
        raise ValueError("Hay filas duplicadas (misma fecha y sede)")
    return df


def calcular(df):
    df = df.assign(deficit_kwh=(df["demanda_kwh"] - df["generacion_kwh"]).clip(lower=0))
    kpis = {
        "energia_total_kwh": round(float(df["generacion_kwh"].sum()), 2),
        "demanda_total_kwh": round(float(df["demanda_kwh"].sum()), 2),
        "deficit_total_kwh": round(float(df["deficit_kwh"].sum()), 2),
    }
    por_sede = (df.groupby(["sede_id", "region"], as_index=False)
                    .agg(generacion=("generacion_kwh", "sum"), demanda=("demanda_kwh", "sum"),
                    deficit=("deficit_kwh", "sum")))
    por_sede["pct_del_deficit"] = por_sede["deficit"] / kpis["deficit_total_kwh"] * 100
    por_sede["pct_de_la_demanda"] = por_sede["demanda"] / kpis["demanda_total_kwh"] * 100
    por_sede["pct_propio_sin_cubrir"] = por_sede["deficit"] / por_sede["demanda"] * 100
    return kpis, por_sede.sort_values("deficit", ascending=False).reset_index(drop=True)


def fmt(x):
    return f"{x:,.0f}"


def textos(df, kpis, por_sede):
    s = por_sede.iloc[0]
    n_dias = df["fecha"].nunique()
    cobertura = 100 * (1 - kpis["deficit_total_kwh"] / kpis["demanda_total_kwh"]) 
    hallazgo = (f"La sede {s['sede_id']} ({s['region']}) concentra {s['pct_del_deficit']:.1f}% del déficit total "
                f"({fmt(s['deficit'])} kWh) con solo {s['pct_de_la_demanda']:.1f}% de la demanda; "
                f"deja sin cubrir {s['pct_propio_sin_cubrir']:.1f}% de su propia demanda. "
                f"La cobertura global de la demanda es {cobertura:.2f}% (100 x (1 - déficit / demanda)).")
    recomendacion = (f"Priorizar la sede {s['sede_id']} para ampliar generación o respaldo energético y "
                    f"revisar también {por_sede.iloc[1]['sede_id']} (segundo mayor déficit); "
                    f"meta: reducir su déficit y darle seguimiento semanal con el KPI de déficit.")
    limitacion = (f"El análisis cubre {n_dias} días ({df['fecha'].min():%d/%m/%Y} a {df['fecha'].max():%d/%m/%Y}) y solo "
                    f"energía en kWh: no incluye capacidad instalada, costos ni almacenamiento, y el déficit no compensa "
                    f"excedentes de unas sedes con faltantes de otras.")
    return hallazgo, recomendacion, limitacion


def grafico(por_sede, ruta_png):
    d = por_sede.sort_values("sede_id")
    x = range(len(d))
    fig, ax = plt.subplots(figsize=(10, 5))
    ax.bar([i - 0.2 for i in x], d["generacion"] / 1000, width=0.4, label="Generación", color="#2e7d32")
    ax.bar([i + 0.2 for i in x], d["demanda"] / 1000, width=0.4, label="Demanda", color="#c0392b")
    top = por_sede.iloc[0]["sede_id"]
    for i, (sede, deficit) in enumerate(zip(d["sede_id"], d["deficit"])):
        ax.text(i, max(d["demanda"].max(), d["generacion"].max()) / 1000 * 1.03, f"{deficit/1000:,.1f}",
                ha="center", fontsize=9, fontweight="bold" if sede == top else "normal",
                color="#c0392b" if sede == top else "#555")
    ax.set_xticks(list(x)); ax.set_xticklabels(d["sede_id"])
    ax.set_ylabel("Energía acumulada (MWh)"); ax.set_xlabel("Sede")
    ax.set_title("Generación vs. demanda por sede (número sobre cada sede = déficit en MWh)", fontsize=11)
    ax.spines[["top", "right"]].set_visible(False); ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.14), ncol=2, frameon=False)
    fig.tight_layout(); fig.savefig(ruta_png, dpi=150); plt.close(fig)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", default=str(CSV_DEFECTO), help="ruta de operacion_sedes.csv")
    ap.add_argument("--salida", default=str(AQUI), help="carpeta donde se escriben los productos")
    a = ap.parse_args()
    ruta_csv, salida = Path(a.csv), Path(a.salida)
    salida.mkdir(parents=True, exist_ok=True)

    df = cargar(ruta_csv)
    kpis, por_sede = calcular(df)
    hallazgo, recomendacion, limitacion = textos(df, kpis, por_sede)
    periodo = f"{df['fecha'].min():%d/%m/%Y} al {df['fecha'].max():%d/%m/%Y}"
    ahora = datetime.now()

    # 1) grafico.png
    png = salida / "grafico.png"
    grafico(por_sede, png)

    # 2) reporte_kpis.csv 
    pd.DataFrame({"indicador": list(kpis), "valor": list(kpis.values())}).to_csv(
        salida / "reporte_kpis.csv", index=False)

    # 3) metadatos.json 
    sha = sha256_archivo(ruta_csv)
    try:
        fuente = ruta_csv.resolve().relative_to(RAIZ).as_posix()
    except ValueError:
        fuente = ruta_csv.name
    meta = {"fuente": fuente, "registros_entrada": int(len(df)), "sha256_fuente": sha,
            "periodo": periodo, "definicion_deficit": DEFINICION_DEFICIT,
            "kpis": kpis, "generado_en": ahora.isoformat(timespec="seconds"),
            "generado_por": "generar_reporte.py"}
    (salida / "metadatos.json").write_text(json.dumps(meta, indent=2, ensure_ascii=False), encoding="utf-8")

    # 4) informe.html
    titulo = "Informe gerencial de energía por sedes"
    tarjetas = [("Energía generada", fmt(kpis["energia_total_kwh"]), "kWh"),
                ("Demanda", fmt(kpis["demanda_total_kwh"]), "kWh"),
                ("Déficit", fmt(kpis["deficit_total_kwh"]), "kWh")]
    img64 = "data:image/png;base64," + base64.b64encode(png.read_bytes()).decode()
    env = Environment(loader=FileSystemLoader(str(AQUI)), autoescape=True)
    html = env.get_template("plantilla.html").render(
        titulo=titulo, periodo=periodo, fecha_generacion=ahora.strftime("%d/%m/%Y %H:%M"), tarjetas=tarjetas,
        grafico=img64, hallazgo=hallazgo, recomendacion=recomendacion, limitacion=limitacion,
        fuente=fuente, registros=len(df), sha256=sha, definicion=DEFINICION_DEFICIT)
    (salida / "informe.html").write_text(html, encoding="utf-8")

    # 5) informe.docx
    doc = Document()
    doc.add_heading(titulo, 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Periodo: {periodo}  |  Generado: {ahora:%d/%m/%Y %H:%M}")
    doc.add_heading("Indicadores clave", 1)
    t = doc.add_table(rows=1, cols=2); t.style = "Light Grid Accent 1"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Indicador", "Valor"
    for nombre, valor, unidad in tarjetas:
        c = t.add_row().cells; c[0].text, c[1].text = nombre, f"{valor} {unidad}"
    doc.add_heading("Gráfico", 1); doc.add_picture(str(png), width=Inches(6))
    for encabezado, texto in (("Hallazgo", hallazgo), ("Recomendación", recomendacion), ("Limitación", limitacion)):
        doc.add_heading(encabezado, 1); doc.add_paragraph(texto)
    p = doc.add_paragraph(f"Fuente: {fuente} | registros: {len(df):,} | SHA-256: {sha}")
    doc.save(str(salida / "informe.docx"))

    print(f"OK | {len(df):,} registros | periodo {periodo}")
    for k, v in kpis.items():
        print(f"   {k}: {v:,.2f}")
    print("Generados en", salida)


if __name__ == "__main__":
    main()
